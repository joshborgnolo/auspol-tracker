#!/usr/bin/env python3
"""Booth-level SEIFA 2021 keying for the 2025 federal election.

Joins every 2025 polling place (AEC GeneralPollingPlacesDownload-31496.csv,
vendored in repo) to its ABS SA1 polygon by point-in-polygon on the AEC
premises coordinates, then attaches all four SEIFA 2021 indexes (IRSD,
IRSAD, IER, IEO — scores + ABS-published national deciles) from the SA1
index cube. Emits data/booth-seifa-2025.json.

Method limits (recorded in _about): the SA1 is the SA1 OF THE BOOTH'S
PREMISES — a geographic proxy for the catchment's SES, not an aggregate of
catchment voters (no booth-catchment SA1 crosswalk exists). AEC coordinates
are ~GDA94/WGS84, polygons are GDA2020; the ~1.8 m datum shift is noise at
SA1 scale. SEIFA 2021 (2021 census) is the newest SEIFA edition; the 2025
election post-dates it by ~3.7 years.

Inputs (provenance-hashed into the output):
  atlas/data/aec-2025/GeneralPollingPlacesDownload-31496.csv   (committed)
  .matilda/seifa-irsad-abs/abs-seifa-2021-sa1-indexes.xlsx     (git-ignored)
  .build/analysis/seat-model/cache/SA1_2021_AUST_SHP_GDA2020.zip (git-ignored)

Re-fetch manifest when a cache is missing:
  SEIFA cube: ABS SEIFA 2021 -> Statistical Area Level 1 (docs/abs-input-sources.md)
  SA1 zip:    https://www.abs.gov.au/statistics/standards/australian-statistical-geography-standard-asgs/edition-3-july-2021-june-2026/access-and-downloads/digital-boundary-files/SA1_2021_AUST_SHP_GDA2020.zip

Deps: python3 + openpyxl. Stdlib for SHP/DBF. Run from repo root.
Guards fail loud. Test: node .build/analysis/seat-model/test-booth-seifa-2025.mjs
"""

import csv
import hashlib
import io
import json
import struct
import sys
import zipfile
from array import array
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
while not (ROOT / ".git").exists():
    ROOT = ROOT.parent

AEC_CSV = ROOT / "atlas/data/aec-2025/GeneralPollingPlacesDownload-31496.csv"
SEIFA_XLSX = ROOT / ".matilda/seifa-irsad-abs/abs-seifa-2021-sa1-indexes.xlsx"
SA1_ZIP = ROOT / ".build/analysis/seat-model/cache/SA1_2021_AUST_SHP_GDA2020.zip"
OUT = ROOT / "data/booth-seifa-2025.json"

AEC_URL = "https://results.aec.gov.au/31496/Website/Downloads/GeneralPollingPlacesDownload-31496.csv"
SA1_URL = ("https://www.abs.gov.au/statistics/standards/"
           "australian-statistical-geography-standard-asgs/edition-3-july-2021-june-2026/"
           "access-and-downloads/digital-boundary-files/SA1_2021_AUST_SHP_GDA2020.zip")

# --- pinned universe facts (verified 2026-10-24 against the vendored CSV) ---
BOOTH_ROWS = 8940
TYPES = {1: 6966, 2: 484, 3: 77, 4: 114, 5: 1299}
STATE_HIST = {"ACT": 147, "NSW": 3028, "NT": 89, "QLD": 1626, "SA": 750,
              "TAS": 329, "VIC": 2072, "WA": 899}
TYPE_NAME = {1: "Polling place", 2: "Special hospital", 3: "Remote mobile",
             4: "Mobile team", 5: "PPVC"}
STATE_ORDER = ["NSW", "VIC", "QLD", "SA", "WA", "TAS", "NT", "ACT"]
STE_OF_DIGIT = {"1": "NSW", "2": "VIC", "3": "QLD", "4": "SA", "5": "WA",
                "6": "TAS", "7": "NT", "8": "ACT", "9": "OT"}

GRID_DEG = 0.25  # grid-cell size for the polygon bbox index


def fail(msg):
    sys.exit(f"GUARD <{msg}>")


def hash16(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


# ---------------------------------------------------------------- AEC booths
def read_booths(path):
    rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig"))))
    hdr = rows[1]
    data = rows[2:]
    ix = {k: hdr.index(k) for k in hdr}
    if len(data) != BOOTH_ROWS:
        fail(f"booth rows {len(data)} != pin {BOOTH_ROWS}")
    thist = Counter(int(r[ix["PollingPlaceTypeID"]]) for r in data)
    if thist != TYPES:
        fail(f"type histogram {dict(thist)} != pin {dict(TYPES)}")
    shist = Counter(r[ix["State"]] for r in data)
    if shist != STATE_HIST:
        fail(f"state histogram {dict(shist)} != pin {dict(STATE_HIST)}")
    booths = []
    seen = set()
    for r in data:
        tid = int(r[ix["PollingPlaceTypeID"]])
        if tid not in TYPE_NAME:
            fail(f"unknown PollingPlaceTypeID {tid}")
        key = (r[ix["DivisionID"]], r[ix["PollingPlaceID"]])
        if key in seen:
            fail(f"dup (div,booth) key {key}")
        seen.add(key)
        lat, lon = r[ix["Latitude"]].strip(), r[ix["Longitude"]].strip()
        lat = float(lat) if lat else None
        lon = float(lon) if lon else None
        booths.append({
            "div": r[ix["DivisionNm"]], "divId": int(r[ix["DivisionID"]]),
            "state": r[ix["State"]], "boothId": int(r[ix["PollingPlaceID"]]),
            "booth": r[ix["PollingPlaceNm"]], "type": TYPE_NAME[tid],
            "venue": r[ix["PremisesNm"]], "suburb": r[ix["PremisesSuburb"]],
            "premisesState": r[ix["PremisesStateAb"]],
            "postcode": r[ix["PremisesPostCode"]],
            "lat": lat, "lon": lon,
        })
    if len({b["divId"] for b in booths}) != 150:
        fail("division count != 150")
    return booths


# ------------------------------------------------------------ SEIFA SA1 cube
def read_seifa(path):
    """Table 1 of the SA1 cube: all four indexes, score + national decile."""
    wb = openpyxl.load_workbook(path, read_only=True)
    sh = next((s for s in wb.worksheets if s.title == "Table 1"),
              max(wb.worksheets, key=lambda s: s.max_row))
    it = sh.iter_rows(values_only=True)
    hdr = None
    for i, row in enumerate(it):
        if i == 5:
            hdr = [str(c).strip() if isinstance(c, str) else c for c in row]
            break
    if hdr is None or hdr[1] != "Score":
        fail("seifa Table 1 header row 5 not found")
    # columns: 0 SA1, 1/2 IRSD score+decile, 3/4 IRSAD, 5/6 IER, 7/8 IEO, 9 URP
    m = {}
    for row in it:
        c0 = row[0]
        if c0 is None or not str(c0).isdigit() or len(str(c0)) != 11:
            continue
        sa1 = str(c0)
        vals = []
        for ci in range(1, 9):
            v = row[ci]
            vals.append(round(float(v), 1) if ci % 2 == 1 and isinstance(v, (int, float))
                        else (int(v) if isinstance(v, (int, float)) else None))
        urp = int(row[9]) if isinstance(row[9], (int, float)) else None
        m[sa1] = tuple(vals) + (urp,)
    if not (59000 <= len(m) <= 60000):
        fail(f"seifa cube rows {len(m)} outside 59k-60k band")
    for sa1, v in m.items():
        for s in (v[0], v[2], v[4], v[6]):
            if s is not None and not (100 <= s <= 1450):
                fail(f"seifa score out of range at {sa1}: {s}")
    print(f"seifa: {len(m)} SA1 rows with scores")
    return m  # sa1 -> (irsd, dec, irsad, dec, ier, dec, ieo, dec, urp)


# ------------------------------------------------------------ SHP + DBF read
def read_dbf_codes(buf):
    nrec, hlen, rlen = struct.unpack_from("<IHH", buf, 4)
    fields = []  # (name, length) — dBase III record starts with delete flag
    off = 32
    pos = 1
    while buf[off] != 0x0D:
        name = buf[off:off + 11].split(b"\x00")[0].decode("ascii")
        flen = buf[off + 16]
        fields.append((name, pos, flen))
        pos += flen
        off += 32
    if "SA1_CODE21" not in [f[0] for f in fields]:
        fail("DBF lacks SA1_CODE21")
    codes = []
    for i in range(nrec):
        rec = buf[hlen + i * rlen: hlen + (i + 1) * rlen]
        for name, rel, flen in fields:
            if name == "SA1_CODE21":
                codes.append(rec[rel:rel + flen].decode("ascii").strip())
                break
    # The ABS SA1 pack ends with one null-geometry sentinel record ZZZZZZZZZZZ
    # ("no spatial" row). Expect exactly one and drop it.
    zeds = [c for c in codes if c == "ZZZZZZZZZZZ"]
    if len(zeds) != 1 or codes[-1] != "ZZZZZZZZZZZ":
        fail("DBF null-record sentinel not as expected")
    codes = [c for c in codes if c != "ZZZZZZZZZZZ"]
    if len(codes) != 61844:
        fail(f"DBF records {len(codes)} != 61844")
    if len(set(codes)) != len(codes):
        fail("DBF SA1 code not unique")
    for c in codes:
        if len(c) != 11 or not c.isdigit() or c[0] not in STE_OF_DIGIT:
            fail(f"DBF odd SA1 code {c!r}")
    return codes


def read_shp_polygons(zbuf):
    """Yield (box, rings) per polygon record; None for the null records."""
    if struct.unpack_from(">i", zbuf, 0)[0] != 9994 or \
       struct.unpack_from("<i", zbuf, 28)[0] != 1000 or \
       struct.unpack_from(">i", zbuf, 24)[0] * 2 != len(zbuf):
        fail("SHP magic bits absent")
    stype = struct.unpack_from("<i", zbuf, 32)[0]
    if stype != 5:
        fail(f"SHP shape type {stype} != 5 (polygon)")
    off = 100
    while off < len(zbuf):
        _rec, clen = struct.unpack_from(">2i", zbuf, off)
        off += 8
        content = zbuf[off:off + clen * 2]
        off += clen * 2
        rtype = struct.unpack_from("<i", content, 0)[0]
        if rtype == 0:
            yield None  # migratory/offshore SA1 or the ZZZ sentinel: no shape
            continue
        if rtype != 5:
            fail("SHP non-polygon record")
        box = struct.unpack_from("<4d", content, 4)
        nparts, npts = struct.unpack_from("<2i", content, 36)
        parts = struct.unpack_from(f"<{nparts}i", content, 44)
        xy = memoryview(content)[44 + 4 * nparts: 44 + 4 * nparts + 16 * npts]
        pts = struct.unpack_from(f"<{npts * 2}d", xy)
        arr = array("d", pts)
        rings = [arr[2 * parts[i]: 2 * (parts[i + 1] if i + 1 < nparts else npts)]
                 for i in range(nparts)]
        yield box, rings


# ------------------------------------------------------------------ PIP join
def ring_crossings(x, y, ring):
    inside = False
    j = len(ring) - 2
    for i in range(0, len(ring), 2):
        xi, yi = ring[i], ring[i + 1]
        xj, yj = ring[j], ring[j + 1]
        j = i
        if (yi > y) == (yj > y):
            continue
        xint = xi + (y - yi) / (yj - yi) * (xj - xi)
        if x < xint:
            inside = not inside
    return inside


def reads_join(sa1zip):
    print("reading SA1 geometry ...")
    z = zipfile.ZipFile(sa1zip)
    codes = read_dbf_codes(z.read("SA1_2021_AUST_GDA2020.dbf"))
    shp = z.read("SA1_2021_AUST_GDA2020.shp")
    polys = []           # per record: (code, bbox, rings)
    null_codes = []
    for code, rec in zip(codes, read_shp_polygons(shp)):
        if rec is None:
            null_codes.append(code)
            continue
        polys.append((code, rec[0], rec[1]))
    # migratory/offshore no-usual-address SA1s carry no polygon (pinned set,
    # observed 2026-10-24: 33 rows — four per state "x979797999x"/"x999949999",
    # except TAS gains 69797979994 and ACT+OT carry only two each)
    want_null = {
        "19797979991", "19797979992", "19797979993", "19999949999",
        "29797979991", "29797979992", "29797979993", "29999949999",
        "39797979991", "39797979992", "39797979993", "39999949999",
        "49797979991", "49797979992", "49797979993", "49999949999",
        "59797979991", "59797979992", "59797979993", "59999949999",
        "69797979991", "69797979992", "69797979993", "69797979994",
        "69999949999",
        "79797979991", "79797979992", "79797979993", "79999949999",
        "89797979991", "89999949999",
        "99797979993", "99999949999",
    }
    if set(null_codes) != want_null:
        fail(f"null-shape SA1s drifted: {sorted(null_codes)[:5]}...")
    if len(polys) != 61811:
        fail(f"SHP polygon records {len(polys)} != 61811")
    # grid index over polygon bboxes
    gs = GRID_DEG
    grid = {}
    for idx, (_c, box, _r) in enumerate(polys):
        x0, y0, x1, y1 = box
        for gx in range(int(x0 / gs), int(x1 / gs) + 1):
            for gy in range(int(y0 / gs), int(y1 / gs) + 1):
                grid.setdefault((gx, gy), []).append(idx)
    def locate(x, y, ring_bfs=2):
        seen, cands = set(), []
        for rad in range(0, ring_bfs + 1):
            gx, gy = int(x / gs), int(y / gs)
            for dx in range(-rad, rad + 1):
                for dy in range(-rad, rad + 1):
                    for i in grid.get((gx + dx, gy + dy), ()):
                        if i not in seen:
                            seen.add(i)
                            box = polys[i][1]
                            if box[0] <= x <= box[2] and box[1] <= y <= box[3]:
                                cands.append(i)
            if cands:
                break
        hits = []
        for i in cands:
            code, _box, rings = polys[i]
            cross = 0
            for ring in rings:
                if ring_crossings(x, y, ring):
                    cross += 1
            if cross % 2 == 1:
                hits.append(code)
        hits.sort()
        return hits[0] if hits else None
    return locate


# --------------------------------------------------------------------- main
def main():
    for p, hint in ((AEC_CSV, "git checkout -- atlas/data/aec-2025"),
                    (SEIFA_XLSX, "docs/abs-input-sources.md refetch"),
                    (SA1_ZIP, f"curl -L {SA1_URL}")):
        if not p.exists():
            fail(f"missing input {p} ({hint})")
    booths = read_booths(AEC_CSV)
    print(f"booths: {len(booths)} rows")
    seifa = read_seifa(SEIFA_XLSX)
    locate = reads_join(SA1_ZIP)

    counts = Counter()
    mism = []
    for b in booths:
        if b["lat"] is None or b["lon"] is None:
            b["note"] = "noCoords"
            counts["noCoords"] += 1
            continue
        if not (-55.0 < b["lat"] < -9.0 and 95.0 < b["lon"] < 169.5):
            b["note"] = "badCoords"
            counts["badCoords"] += 1
            continue
        sa1 = locate(b["lon"], b["lat"])
        if sa1 is None:
            counts["pointOutside"] += 1
            b["note"] = "pointOutside"
            continue
        b["sa1"] = sa1
        ste = STE_OF_DIGIT[sa1[0]]
        if ste == "OT":
            counts["otTerritory"] += 1
        elif ste != b["premisesState"]:
            mism.append((b["div"], b["booth"], sa1, ste, b["premisesState"]))
        rec = seifa.get(sa1)
        if rec is None:
            b["note"] = "seifaSuppressed"
            b["seifa"] = None
            counts["seifaSuppressed"] += 1
        else:
            b["seifa"] = {"irsd": rec[0], "irsad": rec[2], "ier": rec[4],
                          "ieo": rec[6],
                          "dec": {"irsd": rec[1], "irsad": rec[3],
                                  "ier": rec[5], "ieo": rec[7]}}
            counts["scored"] += 1
    counts["mapped_sa1"] = sum(1 for b in booths if b.get("sa1"))
    print("counts:", dict(counts))
    if counts["pointOutside"]:
        todo = [f'{b["div"]}|{b["booth"]}' for b in booths
                if b.get("note") == "pointOutside"]
        print("  pointOutside:", todo)
    if mism:
        print(f"state mismatches ({len(mism)}):")
        for m in mism:
            print("   ", m)
    srate = counts["scored"] / sum(1 for b in booths if b["lat"] is not None
                                   and "pointOutside" not in b.get("note", "")
                                   and b.get("note") != "badCoords")
    print(f"scored rate over coord-holding booths: {srate:.3%}")
    # 2026-10-24 measured 96.78%: the residual is genuine ABS suppression
    # (2,390 of 61,811 SA1s publish no scores; booths in them are low-pop
    # areas). Guard with headroom against that floor, not a fresh fit.
    if srate < 0.96:
        fail("scored rate below 96% (suppression floor is ~3.2%)")

    order = {s: i for i, s in enumerate(STATE_ORDER)}
    booths.sort(key=lambda b: (order[b["state"]], b["div"], b["boothId"]))

    out = {
        "_about": ("W-analysis store: one record per 2025 federal polling "
                   "place keyed to the four SEIFA 2021 indexes (irsd/irsad/"
                   "ier/ieo scores + ABS-published national deciles) via the "
                   "SA1 polygon containing the booth's premises coordinates. "
                   "The SA1 is the SA1 OF THE PREMISES — a geographic proxy "
                   "for the catchment's SES, not an aggregate of catchment "
                   "voters; SEIFA is area-derived, cross-time comparisons "
                   "are rank/decile only, and the 2021 census vintage leads "
                   "the 2025 election by ~3.7 years. Coordinate-less mobile/"
                   "hospital teams carry note=noCoords; rows with (0,0) AEC "
                   "coordinates carry note=badCoords; SA1s score-suppressed "
                   "in the cube carry note=seifaSuppressed with seifa:null. "
                   "stateFlag is only a guard report (premises state vs SA1 "
                   "state); the physical premises location wins."),
        "_provenance": {
            "sourceUrls": {"aec": AEC_URL, "sa1shp": SA1_URL,
                           "seifa": "ABS SEIFA 2021 SA1 cube (see .build/"
                                    "analysis/seifa-index/docs/"
                                    "abs-input-sources.md)"},
            "files": {
                "atlas/data/aec-2025/GeneralPollingPlacesDownload-31496.csv":
                    hash16(AEC_CSV),
                ".matilda/seifa-irsad-abs/abs-seifa-2021-sa1-indexes.xlsx":
                    hash16(SEIFA_XLSX),
                ".build/analysis/seat-model/cache/SA1_2021_AUST_SHP_GDA2020.zip":
                    hash16(SA1_ZIP),
            },
            "extractedAt": str(date.today()),
        },
        "election": 2025,
        "seifaEdition": 2021,
        "geography": "SA1 (ASGS Ed3, GDA2020)",
        "indexes": ["irsd", "irsad", "ier", "ieo"],
        "stateFlagBooths": [f"{d}|{nm}" for d, nm, *_ in mism],
        "counts": dict(counts),
        "booths": [],
    }
    for b in booths:
        rec = {"div": b["div"], "divId": b["divId"], "state": b["state"],
               "boothId": b["boothId"], "booth": b["booth"], "type": b["type"],
               "venue": b["venue"], "suburb": b["suburb"],
               "premisesState": b["premisesState"], "postcode": b["postcode"]}
        if b["lat"] is not None:
            rec["lat"] = round(b["lat"], 6)
            rec["lon"] = round(b["lon"], 6)
        if b.get("note"):
            rec["note"] = b["note"]
        rec["sa1"] = b.get("sa1") or None
        rec["seifa"] = b.get("seifa")  # None when unmappable/suppressed
        out["booths"].append(rec)
    OUT.write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n",
                   encoding="utf-8")
    print(f"OK -> {OUT} ({len(out['booths'])} booths)")


if __name__ == "__main__":
    main()
