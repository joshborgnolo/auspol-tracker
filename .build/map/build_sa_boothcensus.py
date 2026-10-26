#!/usr/bin/env python3
"""Booth-level census lanes (dec/demo) for the sa22/sa26 map frames.

The ECSA venue roll (landed with the booth-coordinate arc) tells the map
WHERE each South Australian polling place stands, so each geocoded booth can
be joined to the census area it stands in — the venue-proxy read the federal
map carries as its booths' `dec` (SEIFA deciles) and `demo` (people cells)
lanes. ECSA releases no votes-by-SA1 product, so the voter-footprint lane
(`vdec`) stays null; only the venue read exists.

  dec  — [irsad, irsd, ier, ieo] NATIONAL deciles of the booth's SA1, taken
         from the ABS SEIFA 2021 cube's own per-area decile columns (the
         column right of each score column, resolved by header title — the
         same titles the district SEIFA pipeline resolves by name; a decile
         is null where the cube publishes no score for the index)
  demo — the SA1's own people cells [lang, a65, a2034, fborn, uni, blue,
         null, null], computed from the 2021 Census GCP SA1 tables with the
         same member and column choices and the same ex-not-stated bases as
         the district people lanes (build_sa26_demo.py), so a booth's row is
         commensurate with its district's. A cell is null where its
         denominator is zero (a tiny SA1 can state no language/occupation).

Inputs (the seat-model/sa26 census pipelines' caches, git-ignored; refetch
manifests live in those pipelines' docstrings):
  .build/analysis/seat-model/cache/gcp-sa1-2021.zip           (GCP 2021 SA1 datapack)
  .build/analysis/seat-model/cache/SA1_2021_AUST_SHP_GDA2020.zip (SA1 2021 polygons)
  .matilda/seifa-irsad-abs/abs-seifa-2021-sa1-indexes.xlsx    (SEIFA 2021 SA1 cube)
  map/data/data-sa22.json / map/data/data-sa26.json           (the map data builds;
        this script splices the two lanes into them and rewrites them — it
        runs after build_sa22_data.py / build_sa26_data.py and is itself rerun
        after them whenever the data builds change, since those wipe the lanes)

Chain:
  python3 -I .build/map/build_sa26_data.py && python3 -I .build/map/build_sa22_data.py
  python3 -I .build/map/build_sa_boothcensus.py
  python3 -I .build/map/build_sa26_past.py map/data && python3 -I .build/map/build_sa22_past.py map/data
  node .build/map/geo_build.mjs .build/map/src/sa26 map/data map/data --years=sa22,sa26
  node .build/map/assemble.mjs

Deps: stdlib + openpyxl (the SEIFA pipelines' cube-reader library).
usage: python3 -I .build/map/build_sa_boothcensus.py
"""
import functools
import json
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sa26_geo import read_sa1_polygons, in_geometry, SA_SENTINELS  # noqa: E402

REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
GCP_ZIP = os.path.join(REPO, ".build", "analysis", "seat-model", "cache", "gcp-sa1-2021.zip")
CUBE_XLSX = os.path.join(REPO, ".matilda", "seifa-irsad-abs", "abs-seifa-2021-sa1-indexes.xlsx")
YEARS = ("sa22", "sa26")

IX = ["irsad", "irsd", "ier", "ieo"]
INDEX_TITLES = {
    "irsad": ("INDEX OF RELATIVE SOCIO-ECONOMIC ADVANTAGE AND DISADVANTAGE",),
    "irsd": ("INDEX OF RELATIVE SOCIO-ECONOMIC DISADVANTAGE",),
    "ier": ("INDEX OF ECONOMIC RESOURCES",),
    "ieo": ("INDEX OF EDUCATION AND OCCUPATION", "INDEX OF EMPLOYMENT AND OCCUPATION"),
}
CUBE_SA_ROWS = 4164  # unknown-suppressed SA1s simply have no row, the district SEIFA pipeline's pin

GCP_ROWS = 61844
GCP_SA_ROWS = 4329
GCP_MEMBER = "2021 Census GCP Statistical Area 1 for AUS/2021Census_G01_AUST_SA1.csv"
DEMO_BANDS = ["0_4", "5_14", "15_19", "20_24", "25_34", "35_44", "45_54",
              "55_64", "65_74", "75_84", "85ov"]
GCP_G60B_OCCS = ["P_Tot_Managers", "P_Tot_Professionals",
                 "P_Tot_TechnicTrades_W", "P_Tot_CommunPersnlSvc_W",
                 "P_Tot_ClericalAdminis_W", "P_Tot_Sales_W",
                 "P_Tot_Mach_oper_drivers", "P_Tot_Labourers"]
GCP_MEMBERS = (
    ("G09F", "2021 Census GCP Statistical Area 1 for AUS/2021Census_G09F_AUST_SA1.csv",
     ["P_Australia_Tot"]),
    ("G09H", "2021 Census GCP Statistical Area 1 for AUS/2021Census_G09H_AUST_SA1.csv",
     ["P_Tot_Tot", "P_COB_NS_Tot"]),
    ("G49B", "2021 Census GCP Statistical Area 1 for AUS/2021Census_G49B_AUST_SA1.csv",
     ["P_PGrad_Deg_Total", "P_GradDip_and_GradCert_Total",
      "P_BachDeg_Total", "P_Lev_Edu_NS_Total"]),
    ("G60B", "2021 Census GCP Statistical Area 1 for AUS/2021Census_G60B_AUST_SA1.csv",
     GCP_G60B_OCCS),
)


def fail(msg):
    raise SystemExit(f"sa-boothcensus: {msg}")


# ------------------------------------------------------------- census tables
@functools.lru_cache(maxsize=None)
def sa1_polygons():
    out = read_sa1_polygons()
    pre = {code: [(min(x for x, _ in o), min(y for _, y in o),
                   max(x for x, _ in o), max(y for _, y in o)) for o, _h in polys]
           for code, polys in out.items()}
    grid = {}
    step = 0.25  # the booth-seifa store's static-grid recipe: polygon bboxes
    # bucketed in 0.25-degree cells, a point's cell plus its expanding rings
    # are its PIP shortlist
    for code, boxes in pre.items():
        for bb in boxes:
            x0 = int(bb[0] / step)
            y0 = int(bb[1] / step)
            x1 = int(bb[2] / step)
            y1 = int(bb[3] / step)
            for i in range(x0, x1 + 1):
                for j in range(y0, y1 + 1):
                    grid.setdefault((i, j), set()).add(code)
    return out, grid, step, pre


def pip(lon, lat):
    sa1, grid, step, boxes = sa1_polygons()
    ci, cj = int(lon / step), int(lat / step)
    # a polygon the point sits in always has grid cells within one ring of
    # the point's cell; the second ring is belt-and-braces for clipped bboxes
    for radius in (0, 1, 2):
        codes = set()
        for i in range(ci - radius, ci + radius + 1):
            for j in range(cj - radius, cj + radius + 1):
                codes |= grid.get((i, j), set())
        hits = [code for code in codes
                if any(bb[0] - 1e-9 <= lon <= bb[2] + 1e-9 and bb[1] - 1e-9 <= lat <= bb[3] + 1e-9
                       for bb in boxes[code])
                and in_geometry((lon, lat), sa1[code])]
        if hits:
            return sorted(hits)[0]
    return None


@functools.lru_cache(maxsize=None)
def seifa_deciles():
    """SA1 code -> {ix: national decile|None} from the ABS 2021 cube's own
    per-area decile columns. Columns resolved by header title (score column),
    the decile column verified to follow it."""
    if not os.path.exists(CUBE_XLSX):
        fail(f"missing SEIFA cube {CUBE_XLSX}")
    wb = openpyxl.load_workbook(CUBE_XLSX, read_only=True)
    sh = wb["Table 1"]
    hdr = list(sh.iter_rows(values_only=True, max_row=6))
    title_row, sub_row = list(hdr[4]), list(hdr[5])
    canon = lambda t: " ".join(str(t).upper().split())
    cols = {}
    for i, t in enumerate(title_row):
        if not t:
            continue
        for ix, wanted in INDEX_TITLES.items():
            if ix in cols:
                continue
            if any(w in canon(t) for w in wanted):
                cols[ix] = i
                break
    if set(cols) != set(IX):
        fail(f"cube header drift: resolved {sorted(cols)}, want {IX}")
    for ix, c in cols.items():
        hist = canon(sub_row[c + 1]) if c + 1 < len(sub_row) and sub_row[c + 1] else ""
        if "DECILE" not in hist:
            fail(f"cube {ix}: column {c + 1} after the score is not a Decile column ({sub_row[c + 1]!r})")
    out, n = {}, 0
    for row in sh.iter_rows(values_only=True, min_row=7):
        v = row[0] if row else None
        if isinstance(v, float):
            v = int(v)
        v = str(v) if v is not None else ""
        if not (len(v) == 11 and v.isdigit() and v.startswith("4")):
            continue
        rec = {}
        for ix, c in cols.items():
            s = row[c]
            d = row[c + 1]
            if not isinstance(s, (int, float)):
                rec[ix] = None
            elif not isinstance(d, (int, float)):
                fail(f"cube {v}: {ix} score {s} with no decile beside it")
            else:
                rec[ix] = int(d)
                if not 1 <= rec[ix] <= 10:
                    fail(f"cube {v}: {ix} decile {rec[ix]} outside 1..10")
        out[v] = rec
        n += 1
    if n != CUBE_SA_ROWS:
        fail(f"cube SA scored rows {n} != pin {CUBE_SA_ROWS}")
    print(f"cube SA deciles: {n} SA1s")
    return out


@functools.lru_cache(maxsize=None)
def demo_cells():
    """SA1 code -> [lang, a65, a2034, fborn, uni, blue, None, None] with the
    district people lanes' exact bases (build_sa26_demo.py shares())."""
    import csv
    import io
    import zipfile
    if not os.path.exists(GCP_ZIP):
        fail(f"missing GCP input {GCP_ZIP}")
    agecols = [f"Age_{b}_yr_P" for b in DEMO_BANDS[:-1]] + ["Age_85ov_P"]
    want = ["SA1_CODE_2021", "Tot_P_P", "Lang_used_home_Eng_only_P",
            "Lang_used_home_Oth_Lang_P", *agecols]
    base = {}
    n = 0
    z = zipfile.ZipFile(GCP_ZIP)
    with io.TextIOWrapper(z.open(GCP_MEMBER), encoding="utf-8-sig") as fh:
        rdr = csv.DictReader(fh)
        miss = [c for c in want if c not in (rdr.fieldnames or [])]
        if miss:
            fail(f"GCP G01 header drift, missing {miss}")
        for row in rdr:
            n += 1
            code = row["SA1_CODE_2021"]
            if not code.startswith("4"):
                continue
            base[code] = (int(row["Tot_P_P"]),
                          int(row["Lang_used_home_Eng_only_P"]),
                          int(row["Lang_used_home_Oth_Lang_P"]),
                          tuple(int(row[c]) for c in agecols))
    if n != GCP_ROWS or len(base) != GCP_SA_ROWS:
        fail(f"GCP G01 rows {n} (pin {GCP_ROWS}), SA {len(base)} (pin {GCP_SA_ROWS})")
    extra = {}
    for tag, member, cols in GCP_MEMBERS:
        if member not in z.namelist():
            fail(f"GCP zip lacks {member}")
        d, m = {}, 0
        with io.TextIOWrapper(z.open(member), encoding="utf-8-sig") as fh:
            rdr = csv.DictReader(fh)
            miss = [c for c in ["SA1_CODE_2021", *cols] if c not in (rdr.fieldnames or [])]
            if miss:
                fail(f"GCP {tag} header drift, missing {miss}")
            for row in rdr:
                m += 1
                code = row["SA1_CODE_2021"]
                if code.startswith("4"):
                    d[code] = tuple(int(row[c]) for c in cols)
        if m != GCP_ROWS or len(d) != GCP_SA_ROWS:
            fail(f"GCP {tag} rows {m}, SA {len(d)}")
        extra[tag] = d
    out = {}
    for code, r in base.items():
        p, eng, oth, bands = r[0], r[1], r[2], r[3]
        aust = extra["G09F"][code][0]
        tot, cobns = extra["G09H"][code]
        uni = sum(extra["G49B"][code][:3])
        eduns = extra["G49B"][code][3]
        oc = extra["G60B"][code]
        blue, occd = oc[2] + oc[6] + oc[7], sum(oc)
        bsum = sum(bands)
        pop15 = p - bands[0] - bands[1]
        def safe(num, den):
            # derive-booth-demo's bounded(): a share of an ex-not-stated
            # basis is none when the basis is empty, AND none when census
            # perturbation of a tiny SA1 takes the arithmetic outside
            # [0,1] (cross-table cells disagree at population ~19) — emit
            # no number rather than clamp one
            if den <= 0:
                return None
            v = num / den
            return round(v, 4) if 0 <= v <= 1 else None
        out[code] = [safe(oth, eng + oth),
                     safe(bands[8] + bands[9] + bands[10], bsum),
                     safe(bands[3] + bands[4], bsum),
                     safe(tot - cobns - aust, tot - cobns),
                     safe(uni, pop15 - eduns),
                     safe(blue, occd),
                     None, None]
    print(f"GCP demo cells: {len(out)} SA1s")
    return out


# ------------------------------------------------------------------ the join
def main():
    decs = seifa_deciles()
    demos = demo_cells()
    for year in YEARS:
        path = os.path.join(REPO, "map", "data", f"data-{year}.json")
        if not os.path.exists(path):
            fail(f"missing {path} — run its data build first")
        with open(path) as f:
            data = json.load(f)
        B = data["booths"]
        n = len(B["id"])
        if not (len(B["lat"]) == n and len(B["lon"]) == n and len(B["ty"]) == n):
            fail(f"{year}: booth lane lengths drifted")
        owned = sum(1 for k in B["dec"] if k is not None)
        if owned:
            fail(f"{year}: {owned} dec rows already filled — the lane is not this script's on a stale data build (rerun build_{year}_data.py first)")
        out_sa1, suppressions = set(), 0
        counts = {"pip": 0, "noscore": 0, "demo": 0, "mobile": 0}
        dec_col, demo_col = [], []
        for k in range(n):
            lat, _lon, ty = B["lat"][k], B["lon"][k], B["ty"][k]
            if ty == 4:
                counts["mobile"] += 1
                dec_col.append(None)
                demo_col.append(None)
                continue
            if lat is None:
                fail(f"{year} booth {B['n'][k]!r}: no coordinates — every fixed booth must carry the venue roll's coords before the census join")
            code = pip(B["lon"][k], lat)
            if code is None:
                fail(f"{year} booth {B['n'][k]!r} at {lat},{B['lon'][k]}: inside no SA1 polygon — an offshore or misplaced venue needs an explicit guard, never a silent miss")
            counts["pip"] += 1
            out_sa1.add(code)
            d = decs.get(code)
            if d is None:
                counts["noscore"] += 1
                dec_col.append(None)
            else:
                if all(v is None for v in d.values()):
                    counts["noscore"] += 1
                suppressions += sum(1 for v in d.values() if v is None)
                dec_col.append([d[ix] for ix in IX])
            dm = demos.get(code)
            if dm is None:
                fail(f"{year} booth {B['n'][k]!r}: SA1 {code} has no GCP row (the GCP and SA1 packs must cover the same codes)")
            counts["demo"] += 1
            demo_col.append(list(dm))
        B["dec"] = dec_col
        B["vdec"] = [None] * n  # ECSA releases no votes-by-SA1 product; the
        # voter-footprint lane is a federal AEC read and stays empty by design
        B["demo"] = demo_col
        with open(path, "w") as f:
            json.dump(data, f, separators=(",", ":"), ensure_ascii=False)
        print(f"{year}: {counts['pip']} booths joined to {len(out_sa1)} SA1s "
              f"({counts['noscore']} in SA1s the cube does not score, "
              f"{suppressions} suppressed index cells, {counts['mobile']} mobile teams null) "
              f"-> {path} ({os.path.getsize(path)} bytes)")


main()
