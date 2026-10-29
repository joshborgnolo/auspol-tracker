"""data/qld<yy>-seifa.json — ABS SEIFA scores per Queensland Legislative
Assembly district, the sa26 discipline on Queensland's own assignments.

The federal seats' SEIFA lane reads the division store's published-or-
derived scores. No publisher derives scores for Queensland's 93 state
districts, so this derives them: the ABS SEIFA SA1 cube's four index scores
(IRSAD, IRSD, IER, IEO) population-weighted over each district's SA1s — the
ABS's own 2039.0 §2.2 recipe the federal division store uses — on the SAME
SA1→district assignment the people-lanes pipeline computed
(.build/map/qld16/assign-sa1.json for qld17's 2016 census read,
.build/map/qld21/assign-sa1.json for qld20 and qld24's 2021 read — a cache
the demo run writes; this script never re-derives the join, so a district's
people figures and its SEIFA scores always cover the same geography).

Scores, not re-compositing: a district's index score is the usual-resident-
population-weighted average of its SA1s' index scores, weighted only where
the cube publishes a score. An SA1 the cube gives no row — or a suppressed
score — drops out of that index only; a district whose joinable cube-
population share pools under W_OBS_MIN is a guard, never an estimate. The
four migratory/offshore sentinel SA1s carry no shape and sit outside the
join, as in the people lanes.

Deciles are ranked among the frame's own 93 district scores (rank_bands on
the weighted scores, equal-count bands 1 = least advantaged), the same
"ranked into tenths of the seats" the federal map's venue-level deciles
carry — not the ABS's national SA1 deciles.

Inputs (caches that already ride this repo's other pipelines):
  .matilda/seifa-irsad-abs/abs-seifa-2021-sa1-indexes.xlsx  (ABS 2033.0.55.001, SEIFA 2021 SA1 Table 1)
  .matilda/seifa-irsad-abs/abs-seifa-2016-sa1-indexes.xls   (the 2016 cube, SSE2/ooxml-legacy xls)
  .build/map/qld{16,21}/assign-sa1.json                     (the demo run's join cache)

Score columns are resolved by header title (the federal reader's rule: the
IRSD/IRSAD lead-column order swaps between editions, a positional read
misreads silently).

Guards fail loud. usage: python3 -I .build/map/build_qld{20,24}_seifa.py
  python3 .build/map/build_qld17_seifa.py   (no -I: xlrd is a user-site
  package of this machine and -I hides user site)
"""
import hashlib
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from qld_geo import QLD_SENTINELS, VINTAGES, load_assign_map, EXPECTED_NAMES

REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))

CUBES = {
    2021: {
        "path": os.path.join(REPO, ".matilda", "seifa-irsad-abs", "abs-seifa-2021-sa1-indexes.xlsx"),
        "kind": "xlsx", "code_len": 11, "rows": 59421, "qld_rows": 12073,
        "product": "ABS 2033.0.55.001 SEIFA 2021, Statistical Area Level 1 indexes (Table 1)",
    },
    2016: {
        "path": os.path.join(REPO, ".matilda", "seifa-irsad-abs", "abs-seifa-2016-sa1-indexes.xls"),
        "kind": "xls", "code_len": 7, "rows": 55140, "qld_rows": 11103,
        "product": "ABS 2033.0.55.001 SEIFA 2016, Statistical Area Level 1 indexes (Table 1)",
    },
}
IX = ["irsad", "irsd", "ier", "ieo"]
INDEX_TITLES = {
    "irsad": ("INDEX OF RELATIVE SOCIO-ECONOMIC ADVANTAGE AND DISADVANTAGE",),
    "irsd": ("INDEX OF RELATIVE SOCIO-ECONOMIC DISADVANTAGE",),
    "ier": ("INDEX OF ECONOMIC RESOURCES",),
    "ieo": ("INDEX OF EDUCATION AND OCCUPATION", "INDEX OF EMPLOYMENT AND OCCUPATION"),
}
W_OBS_MIN = 0.90


def fail(msg):
    raise SystemExit(f"qld-seifa: {msg}")


def hash16(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def _canon(t):
    return " ".join(str(t).upper().split())


def load_cube(year):
    """SA1 code -> {urp, ix: score|None}, QLD rows, columns resolved by
    header title."""
    C = CUBES[year]
    if not os.path.exists(C["path"]):
        fail(f"missing SEIFA cube {C['path']}")
    if C["kind"] == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(C["path"], read_only=True)
        if "Table 1" not in wb.sheetnames:
            fail(f"cube sheets {wb.sheetnames} lack 'Table 1'")
        rows = list(wb["Table 1"].iter_rows(values_only=True))
    else:
        import xlrd
        wb = xlrd.open_workbook(C["path"], on_demand=True)
        sh = wb.sheet_by_name("Table 1")
        rows = [sh.row_values(r) for r in range(sh.nrows)]
    title_i = next(i for i, r in enumerate(rows) if r and any(t and "INDEX OF" in _canon(t) for t in r))
    title_row = [_canon(t) if t is not None else "" for t in rows[title_i]]
    sub_row = rows[title_i + 1] if len(rows) > title_i + 1 else []
    cols = {}
    for ix, wanted in INDEX_TITLES.items():
        for i, t in enumerate(title_row):
            if t and any(w in t for w in wanted):
                cols[ix] = i
                break
        if ix not in cols:
            fail(f"cube header drift: no column titles {wanted!r}")
    urp_col = next((i for i, t in enumerate(rows[title_i])
                    if t and "USUAL RESIDENT POPULATION" in _canon(t)), None)
    if urp_col is None:
        urp_col = next((i for i, t in enumerate(sub_row)
                        if t and "POPULATION" in _canon(t)), None)
    if urp_col is None:
        fail("cube lacks a usual-resident-population column")

    out, nrows = {}, 0
    for row in rows[title_i + 1:]:
        if not row:
            continue
        v = row[0]
        if isinstance(v, float):
            v = int(v)
        s = str(v) if v is not None else ""
        if not (len(s) == C["code_len"] and s.isdigit()):
            continue
        nrows += 1
        if not s.startswith("3"):
            continue
        def cell(i):
            return row[i] if i < len(row) else None
        rec = {"urp": float(cell(urp_col)) if isinstance(cell(urp_col), (int, float)) else 0.0}
        for ix in IX:
            x = cell(cols[ix])
            rec[ix] = float(x) if isinstance(x, (int, float)) else None
        out[s] = rec
    if nrows != C["rows"]:
        fail(f"cube {year} code rows {nrows} != pin {C['rows']}")
    if len(out) != C["qld_rows"]:
        fail(f"cube {year} QLD rows {len(out)} != pin {C['qld_rows']}")
    suppressed = sum(1 for r in out.values() if any(r[ix] is None for ix in IX))
    print(f"cube {year} QLD: {len(out)} rows ({suppressed} carry a suppressed index score)")
    return out


def rank_bands(scores, k=10):
    srt = sorted(scores)
    n = len(srt)
    return [min(k, k * (srt.index(v)) // n + 1) for v in scores]


def seifa_main(frame, year):
    cube = load_cube(year)
    assign = load_assign_map(year)
    order = list(EXPECTED_NAMES)

    joined = [c for c in assign if c in cube]
    unscored = [c for c in assign if c not in cube]
    if unscored:
        print(f"unscored joined SA1s: {len(unscored)} (no cube row; dropped from weightings)")

    pool = {n: {ix: [0.0, 0.0] for ix in IX} for n in order}
    poptot = {n: 0.0 for n in order}
    for code, dist in assign.items():
        r = cube.get(code)
        if r is None:
            continue
        poptot[dist] += r["urp"]
        for ix in IX:
            if r[ix] is not None:
                pool[dist][ix][0] += r[ix] * r["urp"]
                pool[dist][ix][1] += r["urp"]
    cells = {}
    worst = ("", 1.0)
    for dist in order:
        cells[dist] = {}
        for ix in IX:
            sw, w = pool[dist][ix]
            tot = poptot[dist]
            if tot <= 0:
                fail(f"{dist}: no cube population over its SA1s")
            obs = w / tot
            if obs < worst[1]:
                worst = (f"{dist}:{ix}", obs)
            if obs < W_OBS_MIN:
                fail(f"{dist}: {ix} observed across only {100 * obs:.1f}% of joined population — cannot derive")
            cells[dist][ix] = round(sw / w, 1)
    print(f"observed share: minimum {100 * worst[1]:.2f}% joined population ({worst[0]})")

    decs = {ix: rank_bands([cells[n][ix] for n in order]) for ix in IX}
    seifa = {}
    srts = {ix: sorted(cells[n][ix] for n in order) for ix in IX}
    for i, dist in enumerate(order):
        row = {}
        for ix in IX:
            row[ix] = {"v": cells[dist][ix],
                       "pct": round((srts[ix].index(cells[dist][ix]) + 1) / len(order), 4),
                       "dec": decs[ix][i]}
        seifa[dist] = row

    # reconciliation: the district-pool statewide mean must be the cube's own
    # direct QLD mean over joined SA1s (an identity, not a close-enough one)
    for ix in IX:
        num = sum(pool[n][ix][0] for n in order)
        den = sum(pool[n][ix][1] for n in order)
        if den <= 0:
            fail(f"{ix}: zero joined weight")
        statewide = num / den
        direct = [(cube[c][ix], cube[c]["urp"]) for c in joined]
        dnum = sum(s * w for s, w in direct if s is not None)
        dden = sum(w for s, w in direct if s is not None)
        if abs(statewide - dnum / dden) > 1e-6:
            fail(f"{ix}: pooled statewide mean {statewide} != direct {dnum / dden}")
    st = {ix: round(sum(pool[n][ix][0] for n in order) / sum(pool[n][ix][1] for n in order), 1) for ix in IX}

    C = CUBES[year]
    out = {
        "year": frame,
        "jurisdiction": "QLD",
        "censusYear": year,
        "method": ("Population-weighted (ABS 2039.0 §2.2): a district's index score is the "
                   "usual-resident-population-weighted average of its SA1s' cube scores, "
                   "weighted over the SA1s whose score the cube publishes. SA1 assignment is "
                   "the people-lanes join cache, one district per SA1."),
        "weights": "cube usual-resident population",
        "seifa": {n: seifa[n] for n in order},
        "state": st,
        "src": {
            "cube": {"path": os.path.relpath(C["path"], REPO), "sha16": hash16(C["path"]),
                     "product": C["product"]},
            "assignMap": {"path": os.path.relpath(VINTAGES[year]["assign"], REPO),
                          "sha16": hash16(VINTAGES[year]["assign"])},
        },
        "about": (f"ABS SEIFA {year} SA1 index scores (IRSAD, IRSD, IER, IEO) pooled to "
                  f"Queensland's 93 Legislative Assembly districts, usual-resident-population "
                  "weighted per ABS 2039.0 §2.2 — the recipe the federal division store used "
                  "for its own editions, here on the SA1→district assignment the frame's "
                  "people lanes computed. Deciles rank the 93 district scores into tenth-"
                  "equal bands (1 = least advantaged tenth of the 93); pct is the rank "
                  "fraction used for ordering. Suppressed SA1 scores drop out of their "
                  f"index only (observed-population floor {W_OBS_MIN:.0%})."),
    }
    out_path = os.path.join(REPO, "data", f"{frame}-seifa.json")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w") as f:
        json.dump(out, f, separators=(",", ":"), ensure_ascii=False)
    print(f"-> {out_path} ({os.path.getsize(out_path)} bytes); state means {st}")


FRAMES = {"qld24": 2021, "qld20": 2021, "qld17": 2016}
