#!/usr/bin/env python3
"""Booth-level census lanes (dec/demo/cob) for the Queensland map frames
(qld17, qld20, qld24) — the build_sa_boothcensus discipline on the ECQ's
two venue rolls.

The ECQ's venue rolls tell the map WHERE each Queensland polling place
stands (the VRS media feed files lat/lon for every real venue; 2017's day
booths ride the (name, locality) join onto that roll), so each geocoded
booth can be joined to the census area it stands in — the venue-proxy read
the federal map carries as its booths' `dec` (SEIFA deciles, the cube's own
NATIONAL decile column per index) and `demo` (people cells) lanes. ECQ
releases no votes-by-SA1 product, so the voter-footprint lane (`vdec`)
stays null; only the venue read exists.

  dec  — [irsad, irsd, ier, ieo] NATIONAL deciles of the booth's SA1, taken
        from the vintage's ABS SEIFA cube per-area decile columns (score
        column resolved by header title, the decile column verified to
        follow it; a decile is null where the cube publishes no score)
  demo — the SA1's own people cells [lang, a65, a2034, fborn, uni, blue,
        null, null], computed from the G01/G09F/G09G/G09H/uni/occ SA1
        tables with the same member and column choices and the same
        ex-not-stated bases as the district people lanes (the frame's own
        build_qld<yy>_demo.py), so a booth's row is commensurate with its
        district's. A cell is null where its denominator is zero or census
        perturbation takes the arithmetic outside [0,1] (emit no number
        rather than clamp one — the federal bounded() convention).
  cob  — the SA1's ten-country birthplace cell row (COB_KEYS order)

Per-frame vintage: qld17 reads the 2016 census packs, qld20/qld24 the 2021.
Location doctrine per era: on the VRS frames (qld20/qld24) every real booth
must carry coords (the roll files them) — one missing is a guard trip. On
qld17 the (name, locality) coord join stands empty where the era can't
place the venue; those rows' lanes file null, never an estimate. A booth at
a location with NO QLD SA1 polygon (an interstate-hosted centre, an
offshore roving team — pseudo rows never get here) is a guard trip too.

Chain: this script runs after build_qld<yy>_data.py (whose lanes set
dec/demo null) and is itself rerun after any data rebuild, since the data
build wipes the lanes.
usage: python3 -I .build/map/build_qld_boothcensus.py <qld17|qld20|qld24 …>
Deps: stdlib + openpyxl (2021) — and xlrd for qld17's cube (run qld17
without -I: xlrd is a user-site package of this machine).
"""
import functools
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from qld_geo import (QLD_SENTINELS, VINTAGES, in_geometry,  # noqa: E402
                     read_sa1_polygons)
from build_qld_demo_common import COB_KEYS, DEMO_BANDS  # noqa: E402

REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))

IX = ["irsad", "irsd", "ier", "ieo"]
INDEX_TITLES = {
    "irsad": ("INDEX OF RELATIVE SOCIO-ECONOMIC ADVANTAGE AND DISADVANTAGE",),
    "irsd": ("INDEX OF RELATIVE SOCIO-ECONOMIC DISADVANTAGE",),
    "ier": ("INDEX OF ECONOMIC RESOURCES",),
    "ieo": ("INDEX OF EDUCATION AND OCCUPATION", "INDEX OF EMPLOYMENT AND OCCUPATION"),
}
CUBES = {
    2021: os.path.join(REPO, ".matilda", "seifa-irsad-abs", "abs-seifa-2021-sa1-indexes.xlsx"),
    2016: os.path.join(REPO, ".matilda", "seifa-irsad-abs", "abs-seifa-2016-sa1-indexes.xls"),
}
CUBE_KIND = {2021: "xlsx", 2016: "xls"}
CUBE_ROWS = {2021: 59421, 2016: 55140}
CUBE_QLD_ROWS = {2021: 12073, 2016: 11103}

#: booth coordinate policy per frame (see the module docstring)
ERA_UNLOCATED_OK = {"qld17": True, "qld20": False, "qld24": False}
FRAMES = {"qld24": 2021, "qld20": 2021, "qld17": 2016}

UNI_COLS = ["P_PGrad_Deg_Total", "P_GradDip_and_GradCert_Total",
            "P_BachDeg_Total", "P_Lev_Edu_NS_Total"]
OCC_COLS = ["P_Tot_Managers", "P_Tot_Professionals",
            "P_Tot_TechnicTrades_W", "P_Tot_CommunPersnlSvc_W",
            "P_Tot_ClericalAdminis_W", "P_Tot_Sales_W",
            "P_Tot_Mach_oper_drivers", "P_Tot_Labourers"]


def fail(msg):
    raise SystemExit(f"qld-boothcensus: {msg}")


# ------------------------------------------------------------- census tables
@functools.lru_cache(maxsize=None)
def sa1_polygons(year):
    out = read_sa1_polygons(year)
    # bbox per classified poly's outer ring
    pre = {}
    for code, polys in out.items():
        boxes = []
        for outer, _holes in polys:
            xs = [p[0] for p in outer]
            ys = [p[1] for p in outer]
            boxes.append((min(xs), min(ys), max(xs), max(ys)))
        pre[code] = boxes
    grid = {}
    step = 0.25
    for code, boxes in pre.items():
        for bb in boxes:
            x0 = int(bb[0] / step)
            y0 = int(bb[1] / step)
            x1 = int(bb[2] / step)
            y1 = int(bb[3] / step)
            for i in range(x0, x1 + 1):
                for j in range(y0, y1 + 1):
                    grid.setdefault((i, j), set()).add(code)
    return out, grid, step, pre


def pip(year, lon, lat):
    sa1, grid, step, boxes = sa1_polygons(year)
    ci, cj = int(lon / step), int(lat / step)
    for radius in (0, 1, 2):
        codes = set()
        for i in range(ci - radius, ci + radius + 1):
            for j in range(cj - radius, cj + radius + 1):
                codes |= grid.get((i, j), set())
        hits = [code for code in codes
                if any(bb[0] - 1e-9 <= lon <= bb[2] + 1e-9 and bb[1] - 1e-9 <= lat <= bb[3] + 1e-9
                       for bb in boxes[code])
                and in_geometry((lon, lat), sa1[code])]
        if hits:
            return sorted(hits)[0]
    return None


@functools.lru_cache(maxsize=None)
def seifa_deciles(year):
    """SA1 code -> {ix: national decile|None} from that vintage's cube
    (score column resolved by header title, the Decile column verified
    beside it)."""
    import io  # noqa: F401
    path, kind = CUBES[year], CUBE_KIND[year]
    if not os.path.exists(path):
        fail(f"missing SEIFA cube {path}")
    if kind == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        if "Table 1" not in wb.sheetnames:
            fail(f"cube sheets {wb.sheetnames} lack 'Table 1'")
        rows = list(wb["Table 1"].iter_rows(values_only=True))
    else:
        import xlrd
        wb = xlrd.open_workbook(path, on_demand=True)
        sh = wb.sheet_by_name("Table 1")
        rows = [sh.row_values(r) for r in range(sh.nrows)]
    canon = lambda t: " ".join(str(t).upper().split())
    title_i = next(i for i, r in enumerate(rows) if r and any(t and "INDEX OF" in canon(t) for t in r))
    title_row = [canon(t) if t is not None else "" for t in rows[title_i]]
    sub_row = list(rows[title_i + 1]) if len(rows) > title_i + 1 else []
    cols = {}
    for ix, wanted in INDEX_TITLES.items():
        for i, t in enumerate(title_row):
            if t and any(w in t for w in wanted):
                cols[ix] = i
                break
        if ix not in cols:
            fail(f"cube {year} header drift: no column titles {wanted!r}")
    for ix, c in cols.items():
        hist = canon(sub_row[c + 1]) if c + 1 < len(sub_row) and sub_row[c + 1] else ""
        if "DECILE" not in hist:
            fail(f"cube {year} {ix}: column after the score is not a Decile column")
    out, n, nq = {}, 0, 0
    for row in rows[title_i + 1:]:
        if not row:
            continue
        v = row[0]
        if isinstance(v, float):
            v = int(v)
        s = str(v) if v is not None else ""
        if not (len(s) in (7, 11) and s.isdigit()):
            continue
        n += 1
        if not s.startswith("3"):
            continue
        nq += 1
        rec = {}
        for ix, c in cols.items():
            sc = row[c] if c < len(row) else None
            dc = row[c + 1] if c + 1 < len(row) else None
            if not isinstance(sc, (int, float)):
                rec[ix] = None
            elif not isinstance(dc, (int, float)):
                fail(f"cube {year} {s}: {ix} score with no decile beside it")
            else:
                rec[ix] = int(dc)
                if not 1 <= rec[ix] <= 10:
                    fail(f"cube {year} {s}: {ix} decile {rec[ix]} outside 1..10")
        out[s] = rec
    if n != CUBE_ROWS[year] or nq != CUBE_QLD_ROWS[year]:
        fail(f"cube {year} rows {n} (pin {CUBE_ROWS[year]}), QLD {nq} (pin {CUBE_QLD_ROWS[year]})")
    print(f"cube {year} QLD deciles: {nq} SA1s")
    return out


@functools.lru_cache(maxsize=None)
def demo_cells(year):
    """SA1 code -> (demo row, cob row) with the district people lanes'
    exact bases (build_qld_demo_common.py's shares())."""
    import csv
    import io
    import zipfile
    V = VINTAGES[year]
    if not os.path.exists(V["gcp"]):
        fail(f"missing GCP input {V['gcp']}")
    z = zipfile.ZipFile(V["gcp"])
    agecols = [f"Age_{b}_yr_P" for b in DEMO_BANDS[:-1]] + ["Age_85ov_P"]
    want = [V["gcp_key"], "Tot_P_P", *V["lang"], *agecols]
    def member(table):
        name = V["member"].format(table=table)
        if name not in z.namelist():
            fail(f"GCP {year} zip lacks {name}")
        return name
    base = {}
    n = 0
    with io.TextIOWrapper(z.open(member("G01")), encoding="utf-8-sig") as fh:
        rdr = csv.DictReader(fh)
        miss = [c for c in want if c not in (rdr.fieldnames or [])]
        if miss:
            fail(f"GCP {year} G01 header drift, missing {miss}")
        for row in rdr:
            n += 1
            code = row[V["gcp_key"]]
            if not code.startswith("3"):
                continue
            base[code] = (int(row["Tot_P_P"]), int(row[V["lang"][0]]),
                          int(row[V["lang"][1]]),
                          tuple(int(row[c]) for c in agecols))
    extra = {}
    for tag, cols in (("G09F", ["P_Australia_Tot", "P_England_Tot", "P_China_Tot"]),
                      ("G09G", ["P_India_Tot", "P_New_Zealand_Tot", "P_Italy_Tot"]),
                      ("G09H", ["P_Tot_Tot", "P_COB_NS_Tot", "P_Philippines_Tot",
                                "P_Vietnam_Tot", "P_South_Africa_Tot", "P_Scotland_Tot",
                                "P_Sri_Lanka_Tot"]),
                      (V["uni_member"], UNI_COLS),
                      (V["occ_member"], OCC_COLS)):
        d, m = {}, 0
        with io.TextIOWrapper(z.open(member(tag)), encoding="utf-8-sig") as fh:
            rdr = csv.DictReader(fh)
            miss = [c for c in [V["gcp_key"], *cols] if c not in (rdr.fieldnames or [])]
            if miss:
                fail(f"GCP {year} {tag} header drift, missing {miss}")
            for row in rdr:
                m += 1
                code = row[V["gcp_key"]]
                if code.startswith("3"):
                    d[code] = tuple(int(row[c]) for c in cols)
        extra[tag] = d
    out = {}
    for code, r in base.items():
        p, eng, oth, bands = r
        f09 = extra["G09F"][code]
        g09 = extra["G09G"][code]
        h09 = extra["G09H"][code]
        aust, tot, cobns = f09[0], h09[0], h09[1]
        uni = sum(extra[V["uni_member"]][code][:3])
        eduns = extra[V["uni_member"]][code][3]
        oc = extra[V["occ_member"]][code]
        blue, occd = oc[2] + oc[6] + oc[7], sum(oc)
        bsum = sum(bands)
        pop15 = p - bands[0] - bands[1]
        def safe(num, den):
            if den <= 0:
                return None
            v = num / den
            return round(v, 4) if 0 <= v <= 1 else None
        den09 = tot - cobns
        cob = (f09[1], g09[0], f09[2], g09[1], h09[2], h09[3], h09[4],
               g09[2], h09[6], h09[5])
        out[code] = ([safe(oth, eng + oth),
                      safe(bands[8] + bands[9] + bands[10], bsum),
                      safe(bands[3] + bands[4], bsum),
                      safe(tot - cobns - aust, den09),
                      safe(uni, pop15 - eduns),
                      safe(blue, occd),
                      None, None],
                     [safe(c, den09) for c in cob])
    print(f"GCP {year} demo cells: {len(out)} SA1s")
    return out


# ------------------------------------------------------------------ the join
def run_frame(frame):
    year = FRAMES[frame]
    decs = seifa_deciles(year)
    demos = demo_cells(year)
    path = os.path.join(REPO, "map", "data", f"data-{frame}.json")
    if not os.path.exists(path):
        fail(f"missing {path} — run its data build first")
    with open(path) as f:
        data = json.load(f)
    B = data["booths"]
    n = len(B["id"])
    if not (len(B["lat"]) == n and len(B["lon"]) == n and len(B["ty"]) == n):
        fail(f"{frame}: booth lane lengths drifted")
    owned = sum(1 for k in B["dec"] if k is not None)
    if owned:
        fail(f"{frame}: {owned} dec rows already filled — the lane is not"
             " this script's on a stale data build (rerun its data builder first)")
    if data.get("cobKeys"):
        fail(f"{frame}: cobKeys already set — rerun the data builder first")
    unlocated_ok = ERA_UNLOCATED_OK[frame]
    sa1s, suppressions = set(), 0
    counts = {"pip": 0, "noscore": 0, "demo": 0, "mobile": 0, "unlocated": 0}
    dec_col, demo_col, cob_col = [], [], []
    for k in range(n):
        lat, ty = B["lat"][k], B["ty"][k]
        if ty == 4:
            counts["mobile"] += 1
            dec_col.append(None)
            demo_col.append(None)
            cob_col.append(None)
            continue
        if lat is None:
            if not unlocated_ok:
                fail(f"{frame} booth {B['n'][k]!r}: no coordinates — the venue"
                     " roll must carry them before the census join")
            counts["unlocated"] += 1
            dec_col.append(None)
            demo_col.append(None)
            cob_col.append(None)
            continue
        code = pip(year, B["lon"][k], lat)
        if code is None:
            fail(f"{frame} booth {B['n'][k]!r} at {lat},{B['lon'][k]}: inside no"
                 " QLD SA1 polygon — an offshore or interstate venue needs an"
                 " explicit guard, never a silent miss")
        counts["pip"] += 1
        sa1s.add(code)
        d = decs.get(code)
        if d is None or all(v is None for v in d.values()):
            counts["noscore"] += 1
            dec_col.append(None)
        else:
            suppressions += sum(1 for v in d.values() if v is None)
            dec_col.append([d[ix] for ix in IX])
        dm = demos.get(code)
        if dm is None:
            fail(f"{frame} booth {B['n'][k]!r}: SA1 {code} has no GCP row"
                 " (the GCP and SA1 packs must cover the same codes)")
        counts["demo"] += 1
        demo_col.append(list(dm[0]))
        cob_col.append(list(dm[1]))
    B["dec"] = dec_col
    B["vdec"] = [None] * n  # the ECQ releases no votes-by-SA1 product; the
    # voter-footprint lane is a federal AEC read and stays empty by design
    B["demo"] = demo_col
    B["cob"] = cob_col
    data["cobKeys"] = COB_KEYS
    with open(path, "w") as f:
        json.dump(data, f, separators=(",", ":"), ensure_ascii=False)
    print(f"{frame}: {counts['pip']} booths joined to {len(sa1s)} SA1s"
          f" ({counts['noscore']} in SA1s the cube does not score,"
          f" {suppressions} suppressed index cells, {counts['mobile']} mobile"
          f" teams null{', ' + str(counts['unlocated']) + ' unlocated' if counts['unlocated'] else ''})"
          f" -> {path} ({os.path.getsize(path)} bytes)")


def main():
    frames = [a for a in sys.argv[1:] if not a.startswith("-")] or list(FRAMES)
    for frame in frames:
        if frame not in FRAMES:
            fail(f"unknown frame {frame!r}")
    for frame in frames:
        run_frame(frame)


main()
